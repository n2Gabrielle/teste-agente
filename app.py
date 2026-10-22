import sys
import os

# Remove o diretório atual do topo da busca de caminhos do Python
diretorio_atual = os.path.dirname(os.path.abspath(__file__))
if diretorio_atual in sys.path:
    sys.path.remove(diretorio_atual)
sys.path.append(diretorio_atual)

import time
import tracemalloc
import re
import io
import json
import uuid
from datetime import datetime
import requests
import pandas as pd
import gspread
from openpyxl.styles import Alignment, Font, PatternFill
from bs4 import BeautifulSoup
import streamlit as st
from dotenv import load_dotenv

from langchain.agents import AgentExecutor, create_openai_tools_agent
from langchain_core.messages import SystemMessage, HumanMessage
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.callbacks import BaseCallbackHandler
from langchain_community.callbacks.manager import get_openai_callback
from langchain_community.vectorstores import FAISS
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.tools import tool

VECTOR_STORE = None

# Resultados produzidos pelas ferramentas na execução atual.
# As ferramentas de verificação/avaliação e a interface leem daqui, em vez de depender
# do agente repassar (e muitas vezes resumir ou truncar) textos longos.
RESULTADOS = {}

def novo_estado_execucao():
    RESULTADOS.clear()
    RESULTADOS.update({
        "analise_normativa": "",
        "historias": [],
        "lacunas": [],
        "semantica": "",
        "cobertura": "",
        "qus": "",
        "invest": "",
    })

novo_estado_execucao()

def backlog_completo() -> str:
    return "\n\n".join(RESULTADOS["historias"])

def normalizar_id_us(texto: str) -> str:
    return re.sub(r'US-?\s*(\d+)', lambda m: f"US-{int(m.group(1)):03d}", texto, flags=re.IGNORECASE)

def ids_no_backlog() -> list:
    ids = {normalizar_id_us(m) for m in re.findall(r'US-?\s*\d+', backlog_completo(), re.IGNORECASE)}
    return sorted(ids, key=lambda s: int(s.split("-")[1]))

# Configuração da página do Streamlit
st.set_page_config(
    page_title="Sandbox de Engenharia de Requisitos Ágeis - UFF",
    page_icon="🎓",
    layout="wide"
)

# ---------------------------------------------------------
# CONFIGURAÇÃO SEGURA DA API KEY
# ---------------------------------------------------------
load_dotenv(override=True)

api_key = os.getenv("OPENAI_API_KEY")

if not api_key:
    try:
        if "OPENAI_API_KEY" in st.secrets and st.secrets["OPENAI_API_KEY"]:
            api_key = st.secrets["OPENAI_API_KEY"]
    except Exception:
        pass

if not api_key:
    st.error("Chave da OpenAI não configurada. Adicione nos Secrets do Streamlit Cloud ou no arquivo .env local.")
    st.stop()

os.environ["OPENAI_API_KEY"] = api_key

# ---------------------------------------------------------
# EDITOR DE TEMPERATURA (barra lateral)
# ---------------------------------------------------------
TEMPERATURA_PADRAO = 0.0

if "temperatura" not in st.session_state:
    st.session_state["temperatura"] = TEMPERATURA_PADRAO

def restaurar_temperatura():
    st.session_state["temperatura"] = TEMPERATURA_PADRAO

with st.sidebar:
    st.header("⚙️ Configurações do Modelo")
    temperatura = st.slider(
        "🌡️ Temperatura da IA",
        min_value=0.0,
        max_value=2.0,
        step=0.1,
        key="temperatura",
        help="Valores baixos geram respostas mais determinísticas; valores altos, mais criativas e variadas.",
    )
    st.button("↺ Restaurar padrão", on_click=restaurar_temperatura, use_container_width=True)

    if temperatura > 1.0:
        st.warning("Temperaturas acima de 1.0 podem quebrar o formato das seções e da matriz JSON.")
    else:
        st.caption(f"Temperatura atual: **{temperatura:.1f}**")

# Ajustado max_tokens para permitir respostas longas do modelo
llm = ChatOpenAI(model="gpt-4o-mini", temperature=temperatura, openai_api_key=api_key, max_tokens=8012)

def get_response_from_openai(messages):
    return llm.invoke(messages)

def get_text_from_url(url: str) -> str:
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        }
        response = requests.get(url, headers=headers, timeout=15)
        response.raise_for_status()
        
        soup = BeautifulSoup(response.text, 'html.parser')
        
        for element in soup(["script", "style", "nav", "footer", "header"]):
            element.decompose()
            
        text = soup.get_text(separator=' ')
        
        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
        cleaned_text = '\n'.join(chunk for chunk in chunks if chunk)
        
        if not cleaned_text or len(cleaned_text) < 30:
            return f"Aviso: Não foi possível extrair um texto significativo da URL ({url}). Certifique-se de que o documento está publicado na Web."
            
        return cleaned_text
        
    except Exception as e:
        return f"Erro ao extrair conteúdo da URL ({url}): {str(e)}"


def extrair_notas(tabela: str, maximo: int) -> dict:
    """Lê linhas '| US-001 | 6/7 | Atômica |' e devolve {US-001: (6, 'Atômica')}."""
    notas = {}
    padrao = re.compile(r'^\s*\|?\s*\**(US-?\s*\d+)\**\s*\|\s*(\d+)\s*/\s*%d\s*\|\s*([^|]*)' % maximo, re.IGNORECASE)
    for linha in tabela.splitlines():
        m = padrao.match(linha)
        if m:
            notas[normalizar_id_us(m.group(1))] = (int(m.group(2)), m.group(3).strip() or "Nenhuma")
    return notas

def montar_secao_metricas() -> str:
    """Monta a consolidação e a tabela QUS/INVEST em Python a partir das saídas das ferramentas."""
    qus = extrair_notas(RESULTADOS["qus"], 7)
    invest = extrair_notas(RESULTADOS["invest"], 6)
    ids = sorted(set(qus) | set(invest), key=lambda s: int(s.split("-")[1]))

    if not ids:
        return (
            "⚠️ Não foi possível extrair as notas no formato de tabela. Saídas brutas das ferramentas:\n\n"
            f"**QUS**\n\n{RESULTADOS['qus']}\n\n**INVEST**\n\n{RESULTADOS['invest']}"
        )

    total_qus = sum(nota for nota, _ in qus.values())
    total_invest = sum(nota for nota, _ in invest.values())

    linhas = [
        "#### 1. Consolidação Quantitativa das Métricas Sintáticas e Estruturais",
        f"- MÉTRICAS QUS: {total_qus} / {len(qus) * 7}",
        f"- MÉTRICAS INVEST: {total_invest} / {len(invest) * 6}",
    ]

    no_backlog = ids_no_backlog()
    nao_avaliadas = [i for i in no_backlog if i not in qus or i not in invest]
    if nao_avaliadas:
        linhas.append(
            f"\n⚠️ {len(nao_avaliadas)} de {len(no_backlog)} histórias do backlog não receberam nota QUS e/ou INVEST: "
            + ", ".join(nao_avaliadas)
        )

    linhas += [
        "",
        "#### 2. Relatório de Auditoria Estrutural (QUS e INVEST)",
        "",
        "| User Story | Nota QUS | Falhas QUS | Nota INVEST | Falhas INVEST |",
        "| :--- | :--- | :--- | :--- | :--- |",
    ]
    for us in ids:
        nota_q, falha_q = qus.get(us, (None, "-"))
        nota_i, falha_i = invest.get(us, (None, "-"))
        linhas.append(
            f"| {us} | {f'{nota_q}/7' if nota_q is not None else '-'} | {falha_q} "
            f"| {f'{nota_i}/6' if nota_i is not None else '-'} | {falha_i} |"
        )
    return "\n".join(linhas)

# --- EXPORTAÇÃO PARA EXCEL ---
def _limpar_markdown(texto: str) -> str:
    return re.sub(r'[*_`#]+', '', texto).strip().rstrip(',').strip()

def estruturar_backlog() -> list:
    """Converte o texto do backlog em linhas (módulo, ID, título, como, quero, para que, critérios)."""
    padrao_modulo = re.compile(r'^\s*#+\s*Módulo\s*:?\s*(.+)$', re.IGNORECASE)
    padrao_us = re.compile(r'^\W*(US-?\s*\d+)\s*[-–—:]\s*(.+)$', re.IGNORECASE)

    linhas_saida = []
    modulo_atual = ""
    atual = None

    for linha in backlog_completo().splitlines():
        m_mod = padrao_modulo.match(linha)
        m_us = padrao_us.match(linha)
        if m_mod:
            modulo_atual = _limpar_markdown(m_mod.group(1))
        elif m_us:
            atual = {
                "Módulo": modulo_atual,
                "ID": normalizar_id_us(m_us.group(1)),
                "Título": _limpar_markdown(m_us.group(2)),
                "Como (Persona)": "",
                "Quero (Ação)": "",
                "Para que (Benefício)": "",
                "Critérios de Aceitação": [],
            }
            linhas_saida.append(atual)
        elif atual is not None:
            texto = _limpar_markdown(linha.lstrip("-* ").strip())
            if re.match(r'^Como\b', texto, re.IGNORECASE):
                atual["Como (Persona)"] = re.sub(r'^Como\s+', '', texto, flags=re.IGNORECASE)
            elif re.match(r'^quero\b', texto, re.IGNORECASE):
                atual["Quero (Ação)"] = re.sub(r'^quero\s+', '', texto, flags=re.IGNORECASE)
            elif re.match(r'^para que\b', texto, re.IGNORECASE):
                atual["Para que (Benefício)"] = re.sub(r'^para que\s+', '', texto, flags=re.IGNORECASE).rstrip('.')
            elif re.match(r'^CA\s*\d+', texto, re.IGNORECASE):
                atual["Critérios de Aceitação"].append(texto)

    for item in linhas_saida:
        item["Critérios de Aceitação"] = "\n".join(item["Critérios de Aceitação"])
    return linhas_saida

def _formatar_aba(planilha, larguras: dict):
    cabecalho = PatternFill("solid", fgColor="1E293B")
    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cabecalho
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    for linha in planilha.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(vertical="top", wrap_text=True)
    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura
    planilha.freeze_panes = "A2"

def montar_tabelas_resultado(info_execucao: dict):
    """Monta as tabelas (backlog, métricas e consolidação) usadas no .xlsx e no Google Sheets."""
    backlog = estruturar_backlog()
    if not backlog:
        # Formato inesperado: exporta o texto bruto para não perder conteúdo
        backlog = [{"Backlog (texto completo)": backlog_completo()}]
    df_backlog = pd.DataFrame(backlog)

    qus = extrair_notas(RESULTADOS["qus"], 7)
    invest = extrair_notas(RESULTADOS["invest"], 6)
    ids = ids_no_backlog()
    df_metricas = pd.DataFrame([
        {
            "ID": us,
            "Nota QUS (0-7)": qus[us][0] if us in qus else None,
            "Falhas QUS": qus[us][1] if us in qus else "Não avaliada",
            "Nota INVEST (0-6)": invest[us][0] if us in invest else None,
            "Falhas INVEST": invest[us][1] if us in invest else "Não avaliada",
        }
        for us in ids
    ])

    total_qus, max_qus = sum(n for n, _ in qus.values()), len(qus) * 7
    total_invest, max_invest = sum(n for n, _ in invest.values()), len(invest) * 6
    nao_avaliadas = [us for us in ids if us not in qus or us not in invest]
    df_consolidacao = pd.DataFrame(
        [
            ("Data/hora da execução", info_execucao["data_hora"]),
            ("Documento de escopo (URL)", info_execucao["url"]),
            ("Modelo", "gpt-4o-mini"),
            ("Temperatura", info_execucao["temperatura"]),
            ("Histórias no backlog", len(ids)),
            ("MÉTRICAS QUS", f"{total_qus} / {max_qus}"),
            ("QUS (%)", round(100 * total_qus / max_qus, 1) if max_qus else None),
            ("MÉTRICAS INVEST", f"{total_invest} / {max_invest}"),
            ("INVEST (%)", round(100 * total_invest / max_invest, 1) if max_invest else None),
            ("Histórias sem nota QUS e/ou INVEST", ", ".join(nao_avaliadas) or "Nenhuma"),
            ("Tempo total (s)", round(info_execucao["tempo_s"], 2)),
            ("Tokens de entrada", info_execucao["tokens_entrada"]),
            ("Tokens de saída", info_execucao["tokens_saida"]),
            ("Custo estimado (USD)", round(info_execucao["custo_usd"], 5)),
        ],
        columns=["Métrica", "Valor"],
    )
    return df_backlog, df_metricas, df_consolidacao

def gerar_xlsx(info_execucao: dict) -> bytes:
    """Gera a planilha com o backlog RAG ancorado, as métricas QUS/INVEST e a consolidação."""
    df_backlog, df_metricas, df_consolidacao = montar_tabelas_resultado(info_execucao)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_backlog.to_excel(writer, sheet_name="Backlog RAG", index=False)
        df_metricas.to_excel(writer, sheet_name="Métricas QUS-INVEST", index=False)
        df_consolidacao.to_excel(writer, sheet_name="Consolidação", index=False)

        _formatar_aba(writer.sheets["Backlog RAG"], {"A": 28, "B": 10, "C": 38, "D": 28, "E": 50, "F": 50, "G": 90})
        _formatar_aba(writer.sheets["Métricas QUS-INVEST"], {"A": 10, "B": 15, "C": 35, "D": 17, "E": 35})
        _formatar_aba(writer.sheets["Consolidação"], {"A": 38, "B": 70})
    return buffer.getvalue()

# --- HISTÓRICO NO GOOGLE SHEETS ---
# Credencial: conta de serviço do Google, em st.secrets["gcp_service_account"] (Streamlit Cloud)
# ou no arquivo indicado por GCP_SERVICE_ACCOUNT_FILE no .env (execução local).
GOOGLE_SHEET_ID_PADRAO = "19od0s2vx4kGmD0r_2LKxWk7_BHiPRWkqGkezN2OTJmU"
LIMITE_CELULA_SHEETS = 49000  # o Google Sheets aceita até 50.000 caracteres por célula

def conectar_google_sheets():
    credencial = None
    try:
        if "gcp_service_account" in st.secrets:
            credencial = dict(st.secrets["gcp_service_account"])
    except Exception:
        pass

    if credencial:
        cliente = gspread.service_account_from_dict(credencial)
    else:
        caminho = (os.getenv("GCP_SERVICE_ACCOUNT_FILE") or "").strip().strip('"').strip("'")
        if not caminho:
            raise RuntimeError(
                "credencial do Google não configurada: a variável GCP_SERVICE_ACCOUNT_FILE não foi encontrada "
                "no .env (confira se o .env está salvo, na pasta do app.py, e reinicie o Streamlit) "
                "nem há gcp_service_account nos Secrets."
            )
        # Caminho relativo é resolvido a partir da pasta do app.py, não da pasta onde o terminal está
        if not os.path.isabs(caminho):
            caminho = os.path.join(diretorio_atual, caminho)
        if not os.path.exists(caminho):
            arquivos_parecidos = [f for f in os.listdir(diretorio_atual) if f.lower().startswith("credenciais")]
            raise RuntimeError(
                f"arquivo de credencial não encontrado em: {caminho}. "
                f"Arquivos 'credenciais*' na pasta do app: {arquivos_parecidos or 'nenhum'}."
            )
        cliente = gspread.service_account(filename=caminho)

    return cliente.open_by_key(os.getenv("GOOGLE_SHEET_ID") or GOOGLE_SHEET_ID_PADRAO)

def _celula(valor):
    if hasattr(valor, "item"):  # tipos numpy -> tipos Python
        valor = valor.item()
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if isinstance(valor, str) and len(valor) > LIMITE_CELULA_SHEETS:
        return valor[:LIMITE_CELULA_SHEETS] + " …[truncado]"
    return valor

def _obter_aba(planilha, nome: str, cabecalho: list):
    try:
        aba = planilha.worksheet(nome)
    except gspread.WorksheetNotFound:
        aba = planilha.add_worksheet(title=nome, rows=1000, cols=len(cabecalho))
    if not aba.row_values(1):
        aba.update(range_name="A1", values=[cabecalho])
        aba.freeze(rows=1)
        aba.format("1:1", {"textFormat": {"bold": True}})
    return aba

def _status_cobertura() -> str:
    if not RESULTADOS["lacunas"]:
        return "Não verificada"
    if "COBERTURA_TOTAL_CONFIRMADA" in RESULTADOS["lacunas"][-1]:
        return "Confirmada"
    return f"Lacunas encontradas (última verificação de {len(RESULTADOS['lacunas'])})"

def salvar_no_google_sheets(info_execucao: dict) -> str:
    """Acrescenta a execução atual nas abas Execuções, Histórias e Relatórios. Devolve o ID da execução."""
    df_backlog, df_metricas, df_consolidacao = montar_tabelas_resultado(info_execucao)
    id_execucao = f"{datetime.now():%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:4]}"
    data_hora = info_execucao["data_hora"]
    planilha = conectar_google_sheets()

    # Execuções: uma linha por execução, com a consolidação em colunas
    consolidacao = dict(zip(df_consolidacao["Métrica"], df_consolidacao["Valor"]))
    cab_exec = ["ID Execução", *consolidacao.keys(), "Status da cobertura"]
    linha_exec = [id_execucao, *consolidacao.values(), _status_cobertura()]
    _obter_aba(planilha, "Execuções", cab_exec).append_rows(
        [[_celula(v) for v in linha_exec]], value_input_option="RAW"
    )

    # Histórias: uma linha por história, com as notas QUS/INVEST
    df_historias = (
        df_backlog.merge(df_metricas, on="ID", how="left") if "ID" in df_backlog.columns else df_metricas
    )
    cab_hist = ["ID Execução", "Data/hora", *df_historias.columns]
    linhas_hist = [
        [id_execucao, data_hora, *[_celula(v) for v in linha]]
        for linha in df_historias.itertuples(index=False)
    ]
    if linhas_hist:
        _obter_aba(planilha, "Histórias", cab_hist).append_rows(linhas_hist, value_input_option="RAW")

    # Relatórios: textos completos produzidos pelas ferramentas
    cab_rel = ["ID Execução", "Data/hora", "Análise Normativa", "Verificação de Lacunas",
               "Validação Semântica", "Matriz de Cobertura", "Backlog (texto completo)"]
    linha_rel = [
        id_execucao, data_hora,
        RESULTADOS["analise_normativa"],
        "\n\n---\n\n".join(RESULTADOS["lacunas"]),
        RESULTADOS["semantica"],
        RESULTADOS["cobertura"],
        backlog_completo(),
    ]
    _obter_aba(planilha, "Relatórios", cab_rel).append_rows(
        [[_celula(v) for v in linha_rel]], value_input_option="RAW"
    )
    return id_execucao

# --- FERRAMENTAS DO AGENTE ---
@tool
def rag_indexing_tool(url: str) -> str:
    """
    Acessa a URL, realiza o chunking do documento e cria a base vetorial FAISS em memória.
    """
    global VECTOR_STORE
    raw_text = get_text_from_url(url)
    
    if raw_text.startswith("Erro") or raw_text.startswith("Aviso"):
        return raw_text

    # Fragmentação adaptativa otimizada
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=120,
        separators=["\n\n", "\n", ".", ";", " "]
    )
    docs = text_splitter.create_documents([raw_text])
    
    # Geração do Vector Store
    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
    VECTOR_STORE = FAISS.from_documents(docs, embeddings)
    
    return f"Indexação RAG concluída com sucesso. Total de {len(docs)} fragmentos processados."

@tool
def normative_analysis_tool(query_escopo: str) -> str:
    """
    Realiza busca vetorial no documento para identificar lacunas normativas, 
    regras ausentes ou contradições antes de gerar os requisitos.
    """
    global VECTOR_STORE
    if not VECTOR_STORE:
        return "Erro: Base vetorial não inicializada. Execute rag_indexing_tool primeiro."

    docs_relevantes = VECTOR_STORE.similarity_search(query_escopo, k=20)
    contexto_recuperado = "\n---\n".join([d.page_content for d in docs_relevantes])

    messages = [
        SystemMessage(content="""
Você é um Engenheiro de Requisitos especialista em Análise Normativa e Auditoria Primária.
Examine o texto recuperado do documento de escopo e identifique:
1. MAPEAMENTO DE MÓDULOS E REQUISITOS (Liste todos os módulos, APIs, telas, chatbots ou fases citadas).
2. REGRAS OMISSAS OU INCOMPLETAS (Ex: regras de validação ausentes, fluxos alternativos não descritos).
3. DÚVIDAS / LACUNAS que precisam de confirmação do analista.
4. CONTRADIÇÕES internas.

Formato da resposta:
# ANÁLISE NORMATIVA E DIAGNÓSTICO DE LACUNAS
## 1. MÓDULOS E ESCOPO MAPEADOS
- [Módulo/Componente X]: [Resumo das funcionalidades esperadas]
## 2. REGRAS IDENTIFICADAS E VALIDADAS
- [Regras confirmadas no texto]
## 3. LACUNAS E OMISSÕES DETECTADAS
- [Ponto Omisso]: [Explicação da ausência e impacto na regra de negócio]
## 4. PERGUNTAS DE REFINAMENTO SUGERIDAS
- [Pergunta objetiva para o Analista/PO sanar a lacuna]
"""),
        HumanMessage(content=f"Contexto Recuperado via RAG:\n{contexto_recuperado}")
    ]
    resposta = get_response_from_openai(messages).content
    RESULTADOS["analise_normativa"] = resposta
    return resposta

@tool
def user_story_rag_generation_tool(payload_contexto: str) -> str:
    """
    Gera um Backlog de Histórias de Usuário exaustivo e atômico,
    garantindo cobertura completa de todos os módulos e regras do documento.
    As histórias geradas são armazenadas automaticamente no backlog compartilhado.
    Em uma segunda chamada, envie em 'payload_contexto' apenas a lista de pendências.
    """
    global VECTOR_STORE
    if not VECTOR_STORE:
        return "Erro: Base vetorial não inicializada."

    docs_relevantes = VECTOR_STORE.similarity_search(
        "Módulos, funcionalidades, APIs, Chatbot, Heliasta, Sistema Jurídico Web, regras de negócio e estimativas do projeto", 
        k=25
    )
    contexto_completo = "\n---\n".join([f"[Trecho {i+1}]: {d.page_content}" for i, d in enumerate(docs_relevantes)])

    # Na rodada de complemento (após a verificação de lacunas), evita repetir histórias e IDs
    instrucao_continuidade = ""
    if RESULTADOS["historias"]:
        existentes = ids_no_backlog()
        proximo = int(existentes[-1].split("-")[1]) + 1 if existentes else 1
        titulos = "\n".join(re.findall(r'US-?\s*\d+\s*-\s*[^\n*]+', backlog_completo(), re.IGNORECASE))
        instrucao_continuidade = (
            f"\n\nHISTÓRIAS JÁ EXISTENTES NO BACKLOG (NÃO repita):\n{titulos}\n\n"
            f"Gere SOMENTE as histórias faltantes, numerando a partir de US-{proximo:03d}."
        )

    prompt_mapeamento = [
        SystemMessage(content="Com base no contexto fornecido, liste TODOS os módulos, APIs, telas e componentes funcionais individualmente que necessitam de User Stories. Seja exaustivo e liste todos sem exceção."),
        HumanMessage(content=contexto_completo)
    ]
    mapa_funcionalidades = get_response_from_openai(prompt_mapeamento).content

    prompt_geracao = [
        SystemMessage(content="""
Você é um Engenheiro de Requisitos Sênior especializado em detalhamento fino de software industrial.
Sua missão é gerar um Backlog de Histórias de Usuário EXAUSTIVO e COMPLETO cobrindo TODOS os módulos identificados no escopo.

EXIGÊNCIAS DE COBERTURA:
- Você DEVE gerar User Stories para TODOS os módulos e componentes mapeados no documento sem omissões.
- Crie quantas histórias (US-001, US-002, ..., US-N) forem necessárias para cobrir 100% do escopo.

EVITE GENERALISMOS:
- CADA História de Usuário DEVE possuir  Critérios de Aceitação (CA) detalhados que cubram o escopo contendo, por exemplo:
  1. Campos obrigatórios/opcionais envolvidos.
  2. Regras de validação, bloqueios ou pré-condições.
  3. Formatos esperados ou comportamentos do sistema em caso de sucesso/erro.

FORMATO OBRIGATÓRIO PARA CADA HISTÓRIA:

US-[NÚMERO] - [NOME ESPECÍFICO DA FUNCIONALIDADE]
Como <Persona/Ator>,
quero <Ação Única e Concreta com contexto do negócio>,
para que <Benefício Direto e Métrica/Impacto do negócio>.

Critérios de Aceitação:
- CA1: [Detalhamento concreto da regra, campo, validação ou comportamento] (Origem: Trecho X)
- CA2: [Detalhamento concreto da regra, campo, validação ou comportamento] (Origem: Trecho Y)
"""),
        HumanMessage(content=f"MAPEAMENTO DE MÓDULOS OBRIGATÓRIOS:\n{mapa_funcionalidades}\n\nCONTEXTO DO DOCUMENTO DE ESCOPO:\n{contexto_completo}\n\nORIENTAÇÃO ADICIONAL: {payload_contexto}{instrucao_continuidade}")
    ]
    resposta = get_response_from_openai(prompt_geracao).content
    RESULTADOS["historias"].append(resposta)

    # O agente recebe só um resumo: o conteúdo completo fica no backlog compartilhado,
    # evitando que ele precise copiar (ou resuma/trunque) as histórias nas próximas chamadas.
    ids_gerados = sorted({normalizar_id_us(m) for m in re.findall(r'US-?\s*\d+', resposta, re.IGNORECASE)})
    return (
        f"Histórias geradas e armazenadas no backlog: {len(ids_gerados)} nesta chamada "
        f"({', '.join(ids_gerados)}). Total no backlog: {len(ids_no_backlog())}. "
        "As próximas ferramentas leem o backlog completo automaticamente."
    )

@tool
def gap_coverage_verification_tool(observacao: str = "") -> str:
    """
    Compara a Análise Normativa (escopo total do documento) com as User Stories JÁ GERADAS.
    Identifica seções, APIs, regras de infraestrutura ou módulos do documento que NÃO possuem histórias associadas.
    Retorna a lista exata do delta (requisitos faltantes) ou 'COBERTURA_TOTAL_CONFIRMADA'.
    A análise normativa e o backlog completo são lidos automaticamente; NÃO copie o conteúdo no argumento.
    """
    payload = (
        f"ANÁLISE NORMATIVA (ESCOPO TOTAL):\n{RESULTADOS['analise_normativa']}\n\n"
        f"BACKLOG DE USER STORIES GERADO ATÉ AGORA:\n{backlog_completo()}"
    )
    messages = [
        SystemMessage(content="""
Você é um Auditor Especialista em Rastreabilidade de Requisitos.
Sua missão é comparar o escopo total extraído do documento com o backlog de User Stories gerado até agora.

Analise atentamente se ficaram de fora:
1. APIs descritas no Roadmap (ex: Criar Caso Não Processual, Agendamentos).
2. Funcionalidades de Infraestrutura/Governança (ex: limitação de tokens LLM, modelos padrão de petição).
3. Recursos de Interface/Sistemas Legados (ex: Login Automático/SSO, visualização de filas no Heliasta).

REGRAS DE RESPOSTA:
- Se houver requisitos/módulos do documento SEM histórias de usuário criadas, liste expressamente cada um deles com detalhes do que deve ser gerado.
- Se 100% dos requisitos do documento estiverem cobertos por histórias, responda ÚNICA e EXCLUSIVAMENTE a frase: "COBERTURA_TOTAL_CONFIRMADA".
"""),
        HumanMessage(content=f"DADOS PARA AUDITORIA DE LACUNAS:\n\n{payload}")
    ]
    resposta = get_response_from_openai(messages).content
    RESULTADOS["lacunas"].append(resposta)
    return resposta

def payload_contexto_e_historias() -> str:
    return (
        f"CONTEXTO (ANÁLISE NORMATIVA DO DOCUMENTO DE ESCOPO):\n{RESULTADOS['analise_normativa']}\n\n"
        f"HISTÓRIAS DE USUÁRIO GERADAS:\n{backlog_completo()}"
    )

@tool
def semantic_consistency_tool(observacao: str = "") -> str:
    """
    Auditoria semântica e verificação de aderência de escopo.
    O contexto e o backlog completo são lidos automaticamente; NÃO copie o conteúdo no argumento.
    """
    payload = payload_contexto_e_historias()
    messages = [
        SystemMessage(content="""
Você é um Auditor Semântico de Requisitos.
Sua função é verificar se CADA User Story PERTENCE ao escopo do projeto.
Uma história pode ser perfeitamente escrita (QUS 7/7), mas ser REPROVADA SEMANTICAMENTE por não constar no escopo.

Para cada User Story responda:
US-XX
Status: [APROVADO SEMANTICAMENTE] ou [REPROVADO SEMANTICAMENTE]
Justificativa: Explicar concisamente se a funcionalidade existe no documento de escopo.
"""),
        HumanMessage(content=f"DADOS PARA ANÁLISE SEMÂNTICA:\n\n{payload}")
    ]
    resposta = get_response_from_openai(messages).content
    RESULTADOS["semantica"] = resposta
    return "Análise semântica concluída e armazenada para exibição."

@tool
def requirements_coverage_tool(observacao: str = "") -> str:
    """
    Matriz de Rastreabilidade entre regras de escopo e User Stories.
    O contexto e o backlog completo são lidos automaticamente; NÃO copie o conteúdo no argumento.
    """
    payload = payload_contexto_e_historias()
    messages = [
        SystemMessage(content="Retorne SOMENTE um JSON válido com a matriz de rastreabilidade entre todas as regras do escopo e as User Stories geradas."),
        HumanMessage(content=f"DADOS PARA RASTREABILIDADE:\n\n{payload}\n\nFormat: [{{\"Regra\":\"...\", \"UserStory\":\"US-XX\", \"Observacao\":\"...\"}}]")
    ]
    resposta = get_response_from_openai(messages).content
    RESULTADOS["cobertura"] = resposta
    return "Matriz de rastreabilidade concluída e armazenada para exibição."

FORMATO_TABELA_AVALIACAO = """

FORMATO OBRIGATÓRIO DA RESPOSTA: somente uma tabela Markdown, uma linha por história, avaliando TODAS as histórias recebidas, sem texto antes ou depois:
| User Story | Nota | Falhas |
| :--- | :--- | :--- |
| US-001 | {exemplo} | {falha_exemplo} |
Use 'Nenhuma' na coluna Falhas quando a nota for máxima."""

@tool
def quality_assessment_tool(observacao: str = "") -> str:
    """
    Avaliação estritamente estrutural/sintática no framework QUS (7 dimensões).
    O backlog completo é lido automaticamente; NÃO copie as histórias no argumento.
    """
    user_stories = backlog_completo()
    messages = [
        SystemMessage(content="""Você é um auditor do framework QUS (Quality of User Stories).
Avalie cada história em 7 critérios binários (1/0):
1. Bem-formada (Well-formed)
2. Atômica (Atomic)
3. Conceitualmente Sólida (Conceptually Sound)
4. Inambígua (Unambiguous)
5. Mínima (Minimal)
6. Sentença Completa (Full Sentence)
7. Estimável (Estimable)

SE A NOTA FOR MENOR QUE 7, VOCÊ É OBRIGADO A NOMEAR EXATAMENTE QUAIS CRITÉRIOS FALHARAM.
Exemplo para nota 6/7: "Falha: Atômica". NUNCA omita a justificativa da perda de ponto.""" + FORMATO_TABELA_AVALIACAO.format(exemplo="6/7", falha_exemplo="Atômica")),
        HumanMessage(content=f"Avalie:\n{user_stories}")
    ]
    resposta = get_response_from_openai(messages).content
    RESULTADOS["qus"] = resposta
    return f"Avaliação QUS concluída para {len(extrair_notas(resposta, 7))} de {len(ids_no_backlog())} histórias."

@tool
def invest_assessment_tool(observacao: str = "") -> str:
    """
    Avaliação de viabilidade ágil no acrônimo INVEST (6 dimensões).
    O backlog completo é lido automaticamente; NÃO copie as histórias no argumento.
    """
    user_stories = backlog_completo()
    messages = [
        SystemMessage(content="""Você é um Agile Coach especialista em INVEST.
Avalie cada história em 6 critérios binários (1/0):
1. Independent
2. Negotiable
3. Valuable
4. Estimable
5. Small
6. Testable

SE A NOTA FOR MENOR QUE 6, VOCÊ É OBRIGADO A NOMEAR EXATAMENTE QUAIS CRITÉRIOS FALHARAM.
Exemplo para nota 5/6: "Falha: Small". NUNCA omita a justificativa da perda de ponto.""" + FORMATO_TABELA_AVALIACAO.format(exemplo="5/6", falha_exemplo="Small")),
        HumanMessage(content=f"Avalie:\n{user_stories}")
    ]
    resposta = get_response_from_openai(messages).content
    RESULTADOS["invest"] = resposta
    return f"Avaliação INVEST concluída para {len(extrair_notas(resposta, 6))} de {len(ids_no_backlog())} histórias."

# --- CONFIGURAÇÃO DO AGENTE E PROMPT ---
# CORREÇÃO: Adicionado 'gap_coverage_verification_tool' à lista de ferramentas
toolkit = [
    rag_indexing_tool,
    normative_analysis_tool,
    user_story_rag_generation_tool,
    gap_coverage_verification_tool,
    semantic_consistency_tool, 
    requirements_coverage_tool, 
    quality_assessment_tool, 
    invest_assessment_tool
]

prompt = ChatPromptTemplate.from_messages([
    ("system", """
Você é um Engenheiro de Software automatizado especialista em Engenharia de Requisitos com suporte a RAG.

IMPORTANTE: as ferramentas compartilham os resultados automaticamente. A análise normativa e TODAS as histórias
geradas ficam armazenadas e são lidas diretamente pelas ferramentas de verificação e avaliação.
NUNCA copie histórias, critérios de aceitação ou relatórios nos argumentos das ferramentas.

Instruções estritas do fluxo com ancoragem e análise de lacunas:

1. Execute 'rag_indexing_tool' enviando a URL informada para vetorizar o documento.
   - SE retornar erro de acesso ao documento, PARE e informe o erro ao usuário.
2. Execute 'normative_analysis_tool' com a busca das principais regras, módulos e personas do projeto para gerar o relatório de lacunas.
3. Execute 'user_story_rag_generation_tool' para gerar o backlog COMPLETO cobrindo TODOS os módulos mapeados na análise normativa.
4. PASSO DE GARANTIA DE COBERTURA 100%:
   - Execute 'gap_coverage_verification_tool'.
   - SE o retorno for 'COBERTURA_TOTAL_CONFIRMADA', avance para o Passo 5.
   - SE a ferramenta listar requisitos/módulos FALTANTES, execute novamente 'user_story_rag_generation_tool' passando em 'payload_contexto' somente a lista de pendências.
   - Faça no máximo UMA rodada de complemento; depois avance para o Passo 5.
5. Os passos a seguir são independentes: chame as quatro ferramentas JUNTAS, na mesma resposta:
   'semantic_consistency_tool', 'requirements_coverage_tool', 'quality_assessment_tool' e 'invest_assessment_tool'.

Ao final, responda APENAS com um resumo curto (no máximo 3 linhas): quantas histórias foram geradas e se a cobertura total foi confirmada.
NÃO reproduza histórias, tabelas ou relatórios na resposta final: a interface exibe os resultados das ferramentas diretamente.
"""),
    ("human", "{input}"),
    MessagesPlaceholder(variable_name="agent_scratchpad"),
])

agent = create_openai_tools_agent(llm, toolkit, prompt)
agent_executor = AgentExecutor(agent=agent, tools=toolkit, verbose=True, max_iterations=15)

# --- ACOMPANHAMENTO EM TEMPO REAL ---
class PipelineProgressHandler(BaseCallbackHandler):
    """
    Exibe na tela cada passo do agente com o tempo gasto:
    - chamadas das ferramentas (entrada e saída resumidas);
    - chamadas do LLM orquestrador entre as ferramentas ("agente decidindo/escrevendo").
    """

    LIMITE_PREVIA = 1500

    def __init__(self, container):
        self.container = container
        self.passo = 0
        self.ferramenta_ativa = None
        self.agente_ativo = None

    def _previa(self, texto) -> str:
        texto = str(texto)
        if len(texto) > self.LIMITE_PREVIA:
            return texto[: self.LIMITE_PREVIA] + f"\n\n… ({len(texto)} caracteres no total)"
        return texto

    def on_llm_start(self, serialized, prompts, **kwargs):
        # Chamadas internas das ferramentas não entram aqui como passo separado
        if self.ferramenta_ativa is None:
            self.passo += 1
            status = self.container.status(f"🧠 {self.passo}. Agente decidindo o próximo passo...", expanded=False)
            self.agente_ativo = (status, time.time(), self.passo)

    on_chat_model_start = on_llm_start

    def on_llm_end(self, response, **kwargs):
        if self.agente_ativo is not None and self.ferramenta_ativa is None:
            status, inicio, passo = self.agente_ativo
            tokens_saida = (response.llm_output or {}).get("token_usage", {}).get("completion_tokens")
            info_tokens = f" · {tokens_saida} tokens escritos" if tokens_saida else ""
            status.update(
                label=f"🧠 {passo}. Agente decidiu o próximo passo — {time.time() - inicio:.1f}s{info_tokens}",
                state="complete",
            )
            self.agente_ativo = None

    def on_tool_start(self, serialized, input_str, **kwargs):
        self.passo += 1
        nome = (serialized or {}).get("name", "ferramenta")
        status = self.container.status(f"⏳ {self.passo}. Executando {nome}...", expanded=True)
        status.caption(f"Entrada enviada pelo agente ({len(str(input_str))} caracteres):")
        status.code(self._previa(input_str), language=None)
        self.ferramenta_ativa = (status, time.time(), self.passo, nome)

    def on_tool_end(self, output, **kwargs):
        if self.ferramenta_ativa is None:
            return
        status, inicio, passo, nome = self.ferramenta_ativa
        status.caption(f"Saída da ferramenta ({len(str(output))} caracteres):")
        status.code(self._previa(output), language=None)
        status.update(label=f"✅ {passo}. {nome} — {time.time() - inicio:.1f}s", state="complete", expanded=False)
        self.ferramenta_ativa = None

    def on_tool_error(self, error, **kwargs):
        if self.ferramenta_ativa is None:
            return
        status, inicio, passo, nome = self.ferramenta_ativa
        status.error(str(error))
        status.update(label=f"❌ {passo}. {nome} — erro após {time.time() - inicio:.1f}s", state="error")
        self.ferramenta_ativa = None

# --- INTERFACE DO STREAMLIT ---

st.title("Geração de Histórias de Usuário - UFF")
st.subheader("Agentes Multi-Ferramentas")
st.markdown("---")

url_padrao = "https://docs.google.com/document/d/e/2PACX-1vTwj4Yh9UVPzqEpHJMprp875O7bW6XRQek_JNl-1ZxriLWvXvWInIxlxaYY4-yTRRTvxNIvUSPkuFbm/pub"
url_documento = st.text_input("Cole aqui a URL pública do documento de escopo (Google Docs publicado na Web ou site):", value=url_padrao)

if st.button("🚀 Iniciar Pipeline de IA", use_container_width=True):
    tracemalloc.start()
    tempo_inicial = time.time()
    
    # Acompanhamento em tempo real: cada ferramenta chamada pelo agente aparece abaixo com entrada e saída
    st.markdown("#### 🛰️ Execução em tempo real")
    st_callback = PipelineProgressHandler(st.container())
    novo_estado_execucao()

    with st.spinner("O Agente está processando o documento, mapeando todos os módulos e gerando o backlog completo... Aguarde."):
        with get_openai_callback() as cb:
            try:
                result = agent_executor.invoke(
                    {
                        "input": f"Execute todo o pipeline de engenharia de requisitos com auditoria semântica, de cobertura, QUS e INVEST para o projeto em: {url_documento}"
                    },
                    config={"callbacks": [st_callback]},
                )
                
                tempo_final = time.time()
                memoria_atual, memoria_pico = tracemalloc.get_traced_memory()
                tracemalloc.stop()
                
                custo_entrada = (cb.prompt_tokens / 1000000) * 0.15
                custo_saida = (cb.completion_tokens / 1000000) * 0.60
                custo_real_calculado = custo_entrada + custo_saida
                
                # Os conteúdos vêm direto das ferramentas (RESULTADOS), não da resposta final do agente
                if not RESULTADOS["historias"]:
                    st.warning(f"O pipeline terminou sem gerar histórias. Resposta do agente:\n\n{result['output']}")
                else:
                    st.success(f"Processamento Concluído com Sucesso! {result['output']}")

                    info_execucao = {
                        "data_hora": datetime.now().strftime("%d/%m/%Y %H:%M"),
                        "url": url_documento,
                        "temperatura": temperatura,
                        "tempo_s": tempo_final - tempo_inicial,
                        "tokens_entrada": cb.prompt_tokens,
                        "tokens_saida": cb.completion_tokens,
                        "custo_usd": custo_real_calculado,
                    }

                    # Histórico automático: falhas aqui não interrompem a exibição dos resultados
                    try:
                        with st.spinner("Salvando o resultado no Google Sheets..."):
                            id_execucao = salvar_no_google_sheets(info_execucao)
                        st.info(f"📚 Resultado salvo no Google Sheets (execução {id_execucao}).")
                    except Exception as erro_sheets:
                        st.warning(f"Não foi possível salvar no Google Sheets: {erro_sheets}")

                    arquivo_xlsx = gerar_xlsx(info_execucao)
                    st.download_button(
                        "📥 Exportar Backlog RAG e Métricas de Qualidade (.xlsx)",
                        data=arquivo_xlsx,
                        file_name=f"backlog_metricas_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        on_click="ignore",  # não reexecuta o app: os resultados continuam na tela após o download
                        use_container_width=True,
                    )

                conteudo_normativa = RESULTADOS["analise_normativa"] or "_A análise normativa não foi executada._"
                conteudo_us = backlog_completo() or "_Nenhuma história gerada._"
                conteudo_metricas = montar_secao_metricas() if (RESULTADOS["qus"] or RESULTADOS["invest"]) else "_As avaliações QUS/INVEST não foram executadas._"
                conteudo_semantica = RESULTADOS["semantica"] or "_A validação semântica não foi executada._"
                conteudo_cobertura = RESULTADOS["cobertura"]

                tab0, tab1, tab2, tab3, tab4 = st.tabs([
                    "📋 Análise Normativa & Lacunas",
                    "📌 Backlog Ancorado (RAG)", 
                    "📊 Métricas de Qualidade (QUS/INVEST)", 
                    "🔍 Validação Semântica & Cobertura", 
                    "⚡ Desempenho & Recursos"
                ])
                
                with tab0:
                    st.header("📋 Análise Normativa e Detecção de Omissões no Documento")
                    st.markdown(conteudo_normativa.strip())
                    for i, relatorio_lacunas in enumerate(RESULTADOS["lacunas"], start=1):
                        with st.expander(f"Verificação de lacunas — rodada {i}"):
                            st.markdown(relatorio_lacunas)
                with tab1:
                    st.header("📌 Backlog de User Stories Gerado")
                    st.caption(f"{len(ids_no_backlog())} histórias no backlog")
                    st.markdown(conteudo_us.strip())
                    
                with tab2:
                    st.header("📊 Métricas e Auditoria Qualitativa (Estrutura e Sintaxe)")
                    st.markdown(conteudo_metricas.strip())
                    
                with tab3:
                    st.header("🔍 Relatório do Avaliador (Aderência de Escopo e Semântica)")
                    st.markdown(
                        """
                        <style>
                        section[data-testid="stMarkdownContainer"] table { width: 100%; }
                        section[data-testid="stMarkdownContainer"] th { background-color: #1E293B; color: white; text-align: center; }
                        section[data-testid="stMarkdownContainer"] td { vertical-align: top; white-space: pre-wrap; }
                        </style>
                        """,
                        unsafe_allow_html=True,
                    )
                    st.markdown("#### 1. Relatório de Validação de Domínio e Nexo Semântico")
                    st.markdown(conteudo_semantica.strip())

                    st.markdown("#### 2. Matriz de Cobertura de Requisitos")
                    if not conteudo_cobertura:
                        st.markdown("_A matriz de cobertura não foi executada._")
                    else:
                        try:
                            json_limpo = re.sub(r'^```(?:json)?\s*|\s*```$', '', conteudo_cobertura.strip())
                            st.dataframe(json.loads(json_limpo), use_container_width=True)
                        except ValueError:
                            st.markdown(conteudo_cobertura)
                    
                with tab4:
                    st.header("⚡ Desempenho Técnico do Pipeline")
                    col1, col2, col3 = st.columns(3)
                    
                    with col1:
                        st.metric(label="⏱ Tempo Total de Execução", value=f"{tempo_final - tempo_inicial:.2f} s")
                        st.metric(label="💾 Consumo de Memória de Pico", value=f"{memoria_pico / (1024 * 1024):.2f} MB")
                    
                    with col2:
                        st.metric(label="📥 Tokens de Entrada (Prompt)", value=f"{cb.prompt_tokens}")
                        st.metric(label="📤 Tokens de Saída (Completion)", value=f"{cb.completion_tokens}")
                    
                    with col3:
                        st.metric(label="🔢 Total de Tokens", value=f"{cb.total_tokens}")
                        st.metric(label="💵 Custo Real (gpt-4o-mini)", value=f"${custo_real_calculado:.5f} USD")
                        st.metric(label="🌡️ Temperatura Utilizada", value=f"{temperatura:.1f}")
                        
            except Exception as e:
                st.error(f"Ocorreu um erro durante a execução do agente: {e}")